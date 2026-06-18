"""Solartron MSTAT workbook -> cycle-level capacity table.

Validated against the real Zenodo LFP dataset (DOI
10.5281/zenodo.22286450, MD5-verified download 2026-10-19).  Layout
of those workbooks: one file per ageing block
(``B{n}_Discharge{x}C_Charge{y}C/Cycle_NNNN.xlsx``), one sheet per
cycle (``..._Step2_CycleNN``) behind a ~124-row text preamble ending
at ``End Header:``, then numeric (time_s, voltage_V, current_A)
rows at ~1 Hz.  Charge current is positive, discharge negative
(verified: charge CC +0.625 A to 3.65 V, discharge -0.625 A to
2.50 V for B1).

Per cycle, charge/discharge capacities are trapezoidal Coulomb
counts split by current sign.  Global cycle numbers are reconstructed
as ``block_end - n_sheets + sheet_position`` (blocks are 50 cycles
early, 100 later; checked: B1 blocks 0050..0400 then 0500..1500 sum
to 1500).  The very first cycle of a cell starts from partial charge,
so its charge capacity/CE are not meaningful; discharge capacity
(after a CV-terminated charge) is.

Workbooks are streamed sheet-by-sheet (read-only), so memory stays
bounded; the full 121-workbook archive processes in ~20-40 min
single-threaded, less with ``--jobs``.
"""
from __future__ import annotations

import io
import re
import zipfile
from concurrent.futures import ProcessPoolExecutor
from pathlib import Path

import numpy as np
import pandas as pd

_FOLDER_RE = re.compile(r"(B\d+)_Discharge([\d.]+)C_Charge([\d.]+)C")
_BLOCK_RE = re.compile(r"Cycle_(\d+)\.xlsx$", re.IGNORECASE)
_SHEET_RE = re.compile(r"_Cycle(\d+)$")


def _coulomb_count(ws) -> dict | None:
    """Integrate one MSTAT cycle sheet; None if no usable data."""
    ts, vs, cs = [], [], []
    started = False
    for row in ws.iter_rows(values_only=True):
        if not started:
            if row[0] and str(row[0]).startswith("End Header"):
                started = True
            continue
        if row[0] is None or row[1] is None or row[2] is None:
            continue
        try:
            ts.append(float(row[0]))
            vs.append(float(row[1]))
            cs.append(float(row[2]))
        except (TypeError, ValueError):
            continue
    if len(ts) < 100:
        return None
    t, v, i = np.asarray(ts), np.asarray(vs), np.asarray(cs)
    if np.any(np.diff(t) <= 0):
        return None
    dt = np.diff(t)
    imid = 0.5 * (i[1:] + i[:-1])
    pos, neg = imid > 0, imid < 0
    q_chg = float(np.sum(imid[pos] * dt[pos]) / 3600.0)
    q_dis = float(-np.sum(imid[neg] * dt[neg]) / 3600.0)
    return {
        "charge_capacity_Ah": q_chg, "discharge_capacity_Ah": q_dis,
        "duration_s": float(t[-1] - t[0]), "n_points": int(len(t)),
        "voltage_min_V": float(v.min()), "voltage_max_V": float(v.max()),
        "mean_discharge_current_A": float(-np.mean(i[i < 0])) if np.any(i < 0) else np.nan,
    }


def parse_workbook_bytes(data: bytes, folder: str, filename: str) -> list[dict]:
    """Parse one MSTAT workbook (bytes) into per-cycle records."""
    from openpyxl import load_workbook

    m = _FOLDER_RE.search(folder)
    b = _BLOCK_RE.search(filename)
    if not m or not b:
        return []
    cell_id = f"{m.group(1)}_Dis{m.group(2)}C_Chg{m.group(3)}C"
    block_end = int(b.group(1))
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheets = [s for s in wb.sheetnames if _SHEET_RE.search(s)]
        sheets.sort(key=lambda s: int(_SHEET_RE.search(s).group(1)))
        records = []
        for pos, name in enumerate(sheets, start=1):
            stats = _coulomb_count(wb[name])
            if stats is None:
                continue
            records.append({
                "cell_id": cell_id,
                "cycle_index": block_end - len(sheets) + pos,
                "temperature_C": 25.0,
                "discharge_crate": float(m.group(2)),
                "charge_crate": float(m.group(3)),
                "depth_of_discharge": 1.0,
                "source_workbook": filename, "source_sheet": name,
                **stats,
            })
        return records
    finally:
        wb.close()


def _workbook_jobs(source: Path) -> list[tuple[str, str, str]]:
    """List (kind, locator, display) for every workbook in a zip or dir."""
    jobs = []
    if source.suffix.lower() == ".zip":
        with zipfile.ZipFile(source) as zf:
            for name in zf.namelist():
                if name.lower().endswith(".xlsx"):
                    jobs.append(("zip", name, name))
    else:
        for p in sorted(source.rglob("*.xlsx")):
            jobs.append(("file", str(p), str(p.relative_to(source))))
    return jobs


def _run_job(args: tuple[str, str, str, str]) -> list[dict]:
    source_kind, locator, display, zip_path = args
    parts = Path(display).parts
    folder = parts[-2] if len(parts) >= 2 else ""
    if source_kind == "zip":
        with zipfile.ZipFile(zip_path) as zf:
            data = zf.read(locator)
    else:
        data = Path(locator).read_bytes()
    return parse_workbook_bytes(data, folder, Path(display).name)


def extract_cycles(source: str | Path, jobs: int = 1) -> pd.DataFrame:
    """Extract the full cycle table from a Zenodo zip or extracted dir."""
    source = Path(source)
    if not source.exists():
        raise FileNotFoundError(source)
    work = [(kind, loc, disp, str(source)) for kind, loc, disp in _workbook_jobs(source)]
    if not work:
        raise ValueError(f"No .xlsx workbooks found under {source}")
    records: list[dict] = []
    if jobs and jobs > 1:
        with ProcessPoolExecutor(max_workers=jobs) as pool:
            for recs in pool.map(_run_job, work, chunksize=1):
                records.extend(recs)
    else:
        for w in work:
            records.extend(_run_job(w))
    df = pd.DataFrame(records)
    return df.sort_values(["cell_id", "cycle_index"]).reset_index(drop=True)
